import os
import openpyxl
from PIL import Image, ImageDraw, ImageFont

TEMPLATE_BJP = r"C:\Users\DELL\Desktop\format\BJP.png"
TEMPLATE_SP = r"C:\Users\DELL\Desktop\format\SP.png"
TEMPLATE_OTHERS = r"C:\Users\DELL\Desktop\format\Others.png"
OUTPUT_DIR = r"C:\Users\DELL\Desktop\Civoratimes\bot\posters"

EXCEL_FILE = r"C:\Users\DELL\Desktop\Civoratimes\bot\UP_2027_Survey_Corrected.xlsx"
wb = openpyxl.load_workbook(EXCEL_FILE)
ws = wb.active

try:
    font_name_bjp = ImageFont.truetype("georgiab.ttf", 55)  
    font_pct_bjp = ImageFont.truetype("arialbd.ttf", 80)   
    font_name_sp = ImageFont.truetype("georgiab.ttf", 50) 
    font_pct_sp = ImageFont.truetype("arialbd.ttf", 70)
except:
    font_name_bjp = ImageFont.load_default()
    font_pct_bjp = ImageFont.load_default()
    font_name_sp = font_name_bjp
    font_pct_sp = font_pct_bjp

COLOR_BLACK = "#111111"
COLOR_BJP_PCT = "#E85D00"
COLOR_SP_PCT = "#C51D25"
COLOR_OTHERS_PCT = "#555555"

CONFIG = {
    'NDA': {
        'template': TEMPLATE_BJP,
        'name_pos': (768, 410),      
        'pct_bjp_pos': (330, 1020),  
        'pct_sp_pos': (768, 1020),   
        'pct_others_pos': (1206, 1020), 
        'font_name': font_name_bjp,
        'font_pct': font_pct_bjp
    },
    'INDIA': {
        'template': TEMPLATE_SP,
        'name_pos': (512, 330),      
        'pct_bjp_pos': (210, 1060),  
        'pct_sp_pos': (512, 1060),
        'pct_others_pos': (814, 1060),
        'font_name': font_name_sp,
        'font_pct': font_pct_sp
    },
    'OTHERS': {
        'template': TEMPLATE_OTHERS,
        'name_pos': (768, 410),      
        'pct_bjp_pos': (330, 1020),  
        'pct_sp_pos': (768, 1020),   
        'pct_others_pos': (1206, 1020), 
        'font_name': font_name_bjp,
        'font_pct': font_pct_bjp
    }
}

def draw_text(draw, text, position, font, text_color):
    x, y = position
    draw.text((x, y), text, font=font, fill=text_color, anchor="mm")

nda_row, india_row, others_row = None, None, None
for row in range(2, ws.max_row+1):
    w = ws.cell(row, 13).value
    if w == 'NDA' and not nda_row: nda_row = row
    if w == 'INDIA' and not india_row: india_row = row
    if w == 'OTHERS' and not others_row: others_row = row
    if nda_row and india_row and others_row: break

# Generate 3 samples first
samples = [
    (nda_row, 'NDA'),
    (india_row, 'INDIA'),
    (others_row, 'OTHERS')
]

for row, winner in samples:
    if not row: continue
    constituency = str(ws.cell(row, 2).value).upper()
    bjp_pct = ws.cell(row, 12).value or 0
    sp_pct = ws.cell(row, 11).value or 0
    others_pct = ws.cell(row, 15).value or 0
    
    bjp_str = f"{bjp_pct:.1f}%"
    sp_str = f"{sp_pct:.1f}%"
    others_str = f"{others_pct:.1f}%"
    
    conf = CONFIG[winner]
    template_path = conf['template']
    
    try:
        img = Image.open(template_path).convert("RGB")
        draw = ImageDraw.Draw(img)
        
        draw_text(draw, constituency, conf['name_pos'], conf['font_name'], COLOR_BLACK)
        draw_text(draw, bjp_str, conf['pct_bjp_pos'], conf['font_pct'], COLOR_BJP_PCT)
        draw_text(draw, sp_str, conf['pct_sp_pos'], conf['font_pct'], COLOR_SP_PCT)
        draw_text(draw, others_str, conf['pct_others_pos'], conf['font_pct'], COLOR_OTHERS_PCT)
        
        safe_name = constituency.replace(" ", "_").replace("/", "_")
        out_path = os.path.join(OUTPUT_DIR, f"SAMPLE_{winner}.png")
        img.save(out_path)
    except Exception as e:
        print(f"Error on {winner}: {e}")
