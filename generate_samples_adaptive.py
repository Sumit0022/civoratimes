import os
import openpyxl
from PIL import Image, ImageDraw, ImageFont

TEMPLATE_BJP = r"C:\Users\DELL\Desktop\format\BJP.png"
TEMPLATE_SP = r"C:\Users\DELL\Desktop\format\SP.png"
TEMPLATE_OTHERS = r"C:\Users\DELL\Desktop\format\Others.png"
OUTPUT_DIR = r"C:\Users\DELL\Desktop\Civoratimes\bot\posters"

EXCEL_FILE = r"C:\Users\DELL\Desktop\Civoratimes\bot\UP_2027_Survey_Corrected.xlsx"
wb = openpyxl.load_workbook(EXCEL_FILE)
ws = wb.active

FONT_NAME_PATH = "georgiab.ttf"
FONT_PCT_PATH = "arialbd.ttf"

COLOR_BLACK = "#111111"
COLOR_BJP_PCT = "#E85D00"
COLOR_SP_PCT = "#C51D25"
COLOR_OTHERS_PCT = "#555555"

# User's perfect coordinates
name_box = [104, 318, 1028, 422]
pct_bjp_box = [144, 964, 361, 1055]
pct_sp_box = [483, 970, 705, 1052]
pct_others_box = [801, 978, 1015, 1063]

CONFIG = {
    'NDA': {
        'template': TEMPLATE_BJP,
        'name_box': name_box,
        'pct_bjp_box': pct_bjp_box,
        'pct_sp_box': pct_sp_box,
        'pct_others_box': pct_others_box,
        'name_max_font': 90, # Set large so short names are big
        'pct_max_font': 70
    },
    'INDIA': {
        'template': TEMPLATE_SP,
        'name_box': name_box,
        'pct_bjp_box': pct_bjp_box,
        'pct_sp_box': pct_sp_box,
        'pct_others_box': pct_others_box,
        'name_max_font': 90,
        'pct_max_font': 70
    },
    'OTHERS': {
        'template': TEMPLATE_OTHERS,
        'name_box': name_box,
        'pct_bjp_box': pct_bjp_box,
        'pct_sp_box': pct_sp_box,
        'pct_others_box': pct_others_box,
        'name_max_font': 90,
        'pct_max_font': 70
    }
}

def draw_adaptive_text(draw, text, bounding_box, font_path, max_font_size, text_color):
    x1, y1, x2, y2 = bounding_box
    max_width = x2 - x1
    max_height = y2 - y1
    center_x = (x1 + x2) // 2
    center_y = (y1 + y2) // 2
    
    font_size = max_font_size
    
    try:
        font = ImageFont.truetype(font_path, font_size)
    except:
        font = ImageFont.load_default()
        draw.text((center_x, center_y), text, font=font, fill=text_color, anchor="mm")
        return

    while font_size > 10:
        bbox = draw.textbbox((0, 0), text, font=font, anchor="mm")
        text_width = bbox[2] - bbox[0]
        text_height = bbox[3] - bbox[1]
        
        if text_width <= max_width and text_height <= max_height:
            break
            
        font_size -= 2
        font = ImageFont.truetype(font_path, font_size)
        
    draw.text((center_x, center_y), text, font=font, fill=text_color, anchor="mm")

nda_row, india_row, others_row = None, None, None
for row in range(2, ws.max_row+1):
    w = ws.cell(row, 13).value
    if w == 'NDA' and not nda_row: nda_row = row
    if w == 'INDIA' and not india_row: india_row = row
    if w == 'OTHERS' and not others_row: others_row = row
    if nda_row and india_row and others_row: break

for row, winner in [(nda_row, 'NDA'), (india_row, 'INDIA'), (others_row, 'OTHERS')]:
    if not row: continue
    
    # Adding ? mark as requested
    constituency = str(ws.cell(row, 2).value).upper() + " ?"
    
    bjp_pct = ws.cell(row, 12).value or 0
    sp_pct = ws.cell(row, 11).value or 0
    others_pct = ws.cell(row, 15).value or 0
    
    conf = CONFIG[winner]
    template_path = conf['template']
    
    try:
        img = Image.open(template_path).convert("RGB")
        draw = ImageDraw.Draw(img)
        
        draw_adaptive_text(draw, constituency, conf['name_box'], FONT_NAME_PATH, conf['name_max_font'], COLOR_BLACK)
        draw_adaptive_text(draw, f'{bjp_pct:.1f}%', conf['pct_bjp_box'], FONT_PCT_PATH, conf['pct_max_font'], COLOR_BJP_PCT)
        draw_adaptive_text(draw, f'{sp_pct:.1f}%', conf['pct_sp_box'], FONT_PCT_PATH, conf['pct_max_font'], COLOR_SP_PCT)
        draw_adaptive_text(draw, f'{others_pct:.1f}%', conf['pct_others_box'], FONT_PCT_PATH, conf['pct_max_font'], COLOR_OTHERS_PCT)
        
        out_path = os.path.join(OUTPUT_DIR, f"SAMPLE_ADAPTIVE_{winner}.png")
        img.save(out_path)
    except Exception as e:
        print(f"Error {winner}: {e}")
